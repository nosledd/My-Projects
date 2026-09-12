"""Native XLSX creation from validated tabular data."""

from __future__ import annotations

from io import BytesIO
from typing import Mapping, Sequence


def create_xlsx(columns: Sequence[str], rows: Sequence[Sequence[object]]) -> bytes:
    """Build a readable one-sheet workbook without touching the filesystem."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as error:
        raise RuntimeError("XLSX support requires openpyxl; install it with pip") from error

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    sheet.append(list(columns))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1455D9")
    for row in rows:
        sheet.append(list(row))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(columns, start=1):
        values = [str(column)] + [str(row[index - 1]) if index <= len(row) else "" for row in rows]
        sheet.column_dimensions[get_column_letter(index)].width = min(max(len(value) for value in values) + 2, 40)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
