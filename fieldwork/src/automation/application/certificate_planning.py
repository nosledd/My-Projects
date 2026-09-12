"""Deterministic planning for PDF certificate batches."""

from __future__ import annotations

from datetime import datetime, timedelta, timezone
from hashlib import sha256
import json
from pathlib import Path
import re
from uuid import uuid4

from automation.domain.models import ActionCategory, ActionPlan, ArtifactReference, PlanStatus, Preview, ProposedAction, UserRequest
from automation.domain.policies import SafetyPolicy


PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z][A-Za-z0-9 _-]*)\s*\}\}")


def is_certificate_request(instruction: str) -> bool:
    return "certificate" in instruction.lower()


def normalise_name(value: str) -> str:
    """Make `Register No` and `REGISTER_NO` compare as the same name."""
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def detect_placeholders(template_path: Path) -> tuple[str, ...]:
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise RuntimeError("Certificate templates require pypdf") from error
    text = "\n".join(page.extract_text() or "" for page in PdfReader(str(template_path)).pages)
    found = tuple(dict.fromkeys(match.group(1).strip() for match in PLACEHOLDER.finditer(text)))
    if not found:
        raise ValueError("The certificate PDF contains no selectable {{PLACEHOLDER}} text.")
    return found


def read_workbook_rows(data_path: Path) -> tuple[tuple[str, ...], list[dict[str, str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("Certificate data requires openpyxl") from error
    workbook = load_workbook(data_path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = tuple(str(value).strip() if value is not None else "" for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
    if not headers or not all(headers):
        raise ValueError("The first Excel row must contain a name for every column.")
    if len({normalise_name(header) for header in headers}) != len(headers):
        raise ValueError("Excel column headings must be unique after normalisation.")
    rows: list[dict[str, str]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = {header: "" if value is None else str(value) for header, value in zip(headers, values)}
        if any(value.strip() for value in row.values()):
            rows.append(row)
    if not rows:
        raise ValueError("The Excel file has no data rows to turn into certificates.")
    return headers, rows


class CertificatePlanner:
    """Validates a PDF/XLSX pair and prepares one confirmation-gated batch."""

    def __init__(self, policy: SafetyPolicy) -> None:
        self._policy = policy

    def plan(self, request: UserRequest, artifacts: tuple[ArtifactReference, ...], store: object) -> tuple[ActionPlan, Preview]:
        pdfs = [item for item in artifacts if item.display_name.lower().endswith(".pdf")]
        workbooks = [item for item in artifacts if item.display_name.lower().endswith(".xlsx")]
        if len(pdfs) != 1 or len(workbooks) != 1 or len(artifacts) != 2:
            raise ValueError("Certificate generation needs exactly one PDF template and one XLSX data file.")
        template, workbook = pdfs[0], workbooks[0]
        placeholders = detect_placeholders(store.path_for(template.artifact_id))
        headers, rows = read_workbook_rows(store.path_for(workbook.artifact_id))
        by_name = {normalise_name(header): header for header in headers}
        mapping = {placeholder: by_name.get(normalise_name(placeholder)) for placeholder in placeholders}
        missing = [placeholder for placeholder, header in mapping.items() if header is None]
        if missing:
            raise ValueError("Missing Excel columns for: " + ", ".join("{{" + item + "}}" for item in missing))
        required_headers = tuple(header for header in mapping.values() if header is not None)
        invalid = [index + 2 for index, row in enumerate(rows) if any(not row[header].strip() for header in required_headers)]
        if invalid:
            raise ValueError("Required certificate values are blank in Excel row(s): " + ", ".join(map(str, invalid[:10])))
        unused = [header for header in headers if header not in required_headers]
        arguments = {
            "format": "certificate_batch",
            "template_artifact_id": template.artifact_id,
            "data_artifact_id": workbook.artifact_id,
            "mapping": mapping,
            "rows": rows,
        }
        action = ProposedAction(uuid4().hex, "certificate.generate_batch", "1.0", ActionCategory.CREATE_OUTPUT, arguments, (template.artifact_id, workbook.artifact_id))
        self._policy.validate(action)
        now = datetime.now(timezone.utc)
        digest = sha256(json.dumps({"request": request.request_id, "args": arguments}, sort_keys=True).encode()).hexdigest()
        plan = ActionPlan(uuid4().hex, request.request_id, digest, (action,), PlanStatus.AWAITING_CONFIRMATION, now, now + timedelta(minutes=15))
        mapping_rows = [["{{" + placeholder + "}}", header] for placeholder, header in mapping.items()]
        warnings = tuple(("Unused Excel column: " + header) for header in unused)
        preview = Preview(plan.plan_id, f"Generate {len(rows)} certificate PDF(s) and one ZIP file", {"columns": ["Certificate placeholder", "Excel column"], "rows": mapping_rows, "record_count": len(rows), "example": rows[0]}, warnings)
        return plan, preview
